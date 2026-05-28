from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def add_page_number(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)

def set_section_page_number_format(section, fmt="decimal", start_at=None):
    sectPr = section._sectPr
    pgNumType = sectPr.find(qn('w:pgNumType'))
    if pgNumType is None:
        pgNumType = OxmlElement('w:pgNumType')
        sectPr.append(pgNumType)
    pgNumType.set(qn('w:fmt'), fmt)
    if start_at is not None:
        pgNumType.set(qn('w:start'), str(start_at))
    else:
        # If no start_at, remove start attribute to make it continuous
        if qn('w:start') in pgNumType.attrib:
            del pgNumType.attrib[qn('w:start')]

def main():
    doc = Document("DTC225210134_Nguyễn Thanh Tuân_CNTTK21CLC.docx")
    
    # 1. Delete all paragraphs from 50 onwards
    for p in list(doc.paragraphs[50:]):
        p._element.getparent().remove(p._element)
        
    print(f"Initial doc sections: {len(doc.sections)}")
    
    # Configure Section 0 (Covers)
    s0 = doc.sections[0]
    s0.page_width = Cm(21)
    s0.page_height = Cm(29.7)
    s0.top_margin = Cm(2)
    s0.bottom_margin = Cm(2)
    s0.left_margin = Cm(3)
    s0.right_margin = Cm(2)
    s0.different_first_page_header_footer = False
    
    # Unlink covers headers/footers
    s0.header.is_linked_to_previous = False
    s0.footer.is_linked_to_previous = False
    # Clear covers headers/footers text
    for p in s0.header.paragraphs:
        p.text = ""
    for p in s0.footer.paragraphs:
        p.text = ""
        
    # 2. Add Section 1 (Front Matter)
    # Note: adding a section automatically splits the document.
    s1 = doc.add_section(WD_SECTION_START.NEW_PAGE)
    s1.page_width = Cm(21)
    s1.page_height = Cm(29.7)
    s1.top_margin = Cm(2)
    s1.bottom_margin = Cm(2)
    s1.left_margin = Cm(3)
    s1.right_margin = Cm(2)
    
    # Configure page numbering for Section 1 (Roman lowercase, continuous from v if needed, or start from i)
    set_section_page_number_format(s1, fmt="romanLowercase", start_at=1)
    s1.footer.is_linked_to_previous = False
    f1_para = s1.footer.paragraphs[0]
    f1_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f1_para.text = ""
    run1 = f1_para.add_run()
    run1.font.name = 'Times New Roman'
    run1.font.size = Pt(12)
    add_page_number(run1)
    
    # Add text in Section 1
    doc.add_heading("PHẦN 1: MẪU TRANG TRƯỚC", level=1)
    doc.add_paragraph("Đây là trang lời cảm ơn...")
    
    # 3. Add Section 2 (LỜI NÓI ĐẦU)
    s2 = doc.add_section(WD_SECTION_START.NEW_PAGE)
    s2.page_width = Cm(21)
    s2.page_height = Cm(29.7)
    s2.top_margin = Cm(2)
    s2.bottom_margin = Cm(2)
    s2.left_margin = Cm(3)
    s2.right_margin = Cm(2)
    
    # Roman lowercase, continuous
    set_section_page_number_format(s2, fmt="romanLowercase")
    s2.footer.is_linked_to_previous = True # inherit from Section 1
    
    # Add text in Section 2
    doc.add_heading("LỜI NÓI ĐẦU", level=1)
    doc.add_paragraph("Lý do chọn đề tài...")
    
    # 4. Add Section 3 (CHƯƠNG TRÌNH CHÍNH)
    s3 = doc.add_section(WD_SECTION_START.NEW_PAGE)
    s3.page_width = Cm(21)
    s3.page_height = Cm(29.7)
    s3.top_margin = Cm(2)
    s3.bottom_margin = Cm(2)
    s3.left_margin = Cm(3)
    s3.right_margin = Cm(2)
    
    # Arabic, start at 1
    set_section_page_number_format(s3, fmt="decimal", start_at=1)
    s3.footer.is_linked_to_previous = False
    f3_para = s3.footer.paragraphs[0]
    f3_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f3_para.text = ""
    run3 = f3_para.add_run()
    run3.font.name = 'Times New Roman'
    run3.font.size = Pt(12)
    add_page_number(run3)
    
    # Add text in Section 3
    doc.add_heading("CHƯƠNG 1: TỔNG QUAN", level=1)
    doc.add_paragraph("Nội dung chương 1...")
    
    doc.save("test_sections_output.docx")
    print(f"Saved! Total sections in output: {len(doc.sections)}")

if __name__ == '__main__':
    main()
