from docx import Document

def main():
    doc = Document("DTC225210134_Nguyễn Thanh Tuân_CNTTK21CLC.docx")
    
    print("SECTION 1 PARAGRAPHS (50 to 79):")
    for idx in range(50, 80):
        if idx < len(doc.paragraphs):
            p = doc.paragraphs[idx]
            print(f"[{idx}]: {repr(p.text)}")
            
    print("\nSECTION 2 PARAGRAPHS (80 to 144):")
    for idx in range(80, 145):
        if idx < len(doc.paragraphs):
            p = doc.paragraphs[idx]
            print(f"[{idx}]: {repr(p.text)}")

if __name__ == '__main__':
    main()
