from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def main():
    doc = Document('DTC225210134_Nguyễn Thanh Tuân_CNTTK21CLC.docx')
    body = doc.element.body
    
    table_index = 0
    paragraph_index = 0
    
    for child in body:
        if child.tag.endswith('p'):
            text = child.text if child.text else ""
            # Let's get the text of the paragraph
            p_obj = [p for p in doc.paragraphs if p._element == child]
            p_text = p_obj[0].text if p_obj else ""
            if p_text.strip():
                print(f"P [{paragraph_index}]: {repr(p_text[:100])}")
            paragraph_index += 1
        elif child.tag.endswith('tbl'):
            t_obj = [t for t in doc.tables if t._element == child]
            t_text = ""
            if t_obj:
                rows = t_obj[0].rows
                if rows:
                    t_text = " | ".join(c.text.strip().replace('\n', ' ')[:30] for c in rows[0].cells)
            print(f"--- TABLE {table_index+1} at body index {body.index(child)}: {t_text} ---")
            table_index += 1

if __name__ == '__main__':
    main()
