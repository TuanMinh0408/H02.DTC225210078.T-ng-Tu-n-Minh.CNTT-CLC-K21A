from docx import Document

def main():
    doc = Document("DTC225210134_Nguyễn Thanh Tuân_CNTTK21CLC.docx")
    print("TOTAL PARAGRAPHS:", len(doc.paragraphs))
    print("TOTAL SECTIONS:", len(doc.sections))
    print("TOTAL TABLES:", len(doc.tables))
    
    with open("tuan_beginning_dump.txt", "w", encoding="utf-8") as f:
        # We want to dump details of all paragraphs up to 250
        for i, p in enumerate(doc.paragraphs[:250]):
            style_name = p.style.name if p.style else "None"
            alignment = p.alignment if p.alignment else "None"
            f.write(f"[{i:03d}] Style: {style_name:<20} | Align: {alignment:<10} | Text: {repr(p.text)}\n")
            
if __name__ == '__main__':
    main()
