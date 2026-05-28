from docx import Document
from docx.oxml.ns import qn

def inspect():
    doc = Document("DTC225210134_Nguyễn Thanh Tuân_CNTTK21CLC.docx")
    for idx, s in enumerate(doc.sections):
        print(f"\n=== Section {idx} ===")
        # Page number properties
        sectPr = s._sectPr
        pgNumType = sectPr.find(qn('w:pgNumType'))
        if pgNumType is not None:
            fmt = pgNumType.get(qn('w:fmt'))
            start = pgNumType.get(qn('w:start'))
            print(f"  Page Numbering: fmt={fmt}, start={start}")
        else:
            print("  Page Numbering: Default")
            
        # Headers
        print(f"  Header different first page: {s.different_first_page_header_footer}")
        for h_type in ['header', 'first_page_header', 'even_page_header']:
            header = getattr(s, h_type, None)
            if header:
                paras = [p.text.strip() for p in header.paragraphs if p.text.strip()]
                print(f"    {h_type} has {len(header.paragraphs)} paragraphs: {paras}")
                
        # Footers
        for f_type in ['footer', 'first_page_footer', 'even_page_footer']:
            footer = getattr(s, f_type, None)
            if footer:
                paras = [p.text.strip() for p in footer.paragraphs if p.text.strip()]
                print(f"    {f_type} has {len(footer.paragraphs)} paragraphs: {paras}")
                
if __name__ == '__main__':
    inspect()
