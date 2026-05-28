from docx import Document

def main():
    doc = Document("DTC225210134_Nguyễn Thanh Tuân_CNTTK21CLC.docx")
    print("TOTAL PARAGRAPHS:", len(doc.paragraphs))
    
    with open("tuan_end_dump.txt", "w", encoding="utf-8") as f:
        for i in range(1580, len(doc.paragraphs)):
            p = doc.paragraphs[i]
            style_name = p.style.name if p.style else "None"
            alignment = p.alignment if p.alignment else "None"
            f.write(f"[{i:04d}] Style: {style_name:<20} | Align: {alignment:<10} | Text: {repr(p.text)}\n")
            
if __name__ == '__main__':
    main()
