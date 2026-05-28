from docx import Document

def main():
    doc = Document('DTC225210078_Tăng Tuấn Minh_CNTTK21CLC.docx')
    print("TOTAL PARAGRAPHS:", len(doc.paragraphs))
    print("TOTAL TABLES:", len(doc.tables))
    for i, table in enumerate(doc.tables):
        rows = len(table.rows)
        cols = len(table.columns)
        print(f"\nTable {i+1}: {rows}x{cols}")
        if rows > 0:
            cells = [cell.text.strip().replace('\n', ' ')[:50] for cell in table.rows[0].cells]
            print(f"   Row 0: {' | '.join(cells)}")
        if rows > 1:
            cells = [cell.text.strip().replace('\n', ' ')[:50] for cell in table.rows[1].cells]
            print(f"   Row 1: {' | '.join(cells)}")

if __name__ == '__main__':
    main()
