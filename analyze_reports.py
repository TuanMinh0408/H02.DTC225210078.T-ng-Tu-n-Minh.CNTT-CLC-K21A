"""
Analyze Reports — Trích xuất và so sánh cấu trúc 2 file docx
================================================================
So sánh cấu trúc báo cáo của Nguyễn Thanh Tuân và Tăng Tuấn Minh.
"""

from docx import Document
from docx.oxml.ns import qn
from pathlib import Path
import json
import re

BASE_DIR = Path(__file__).parent

FILES = {
    'tuan': BASE_DIR / 'DTC225210134_Nguyễn Thanh Tuân_CNTTK21CLC.docx',
    'minh': BASE_DIR / 'DTC225210078_Tăng Tuấn Minh_CNTTK21CLC.docx',
}


def analyze_docx(filepath: str, label: str):
    """Trích xuất toàn bộ cấu trúc từ file docx."""
    print(f"\n{'='*80}")
    print(f"  PHÂN TÍCH: {label}")
    print(f"  File: {Path(filepath).name}")
    print(f"  Size: {Path(filepath).stat().st_size / 1024 / 1024:.1f} MB")
    print(f"{'='*80}")

    doc = Document(filepath)

    # 1. Headings
    headings = []
    para_count = 0
    total_text_len = 0
    for para in doc.paragraphs:
        para_count += 1
        total_text_len += len(para.text)
        style_name = para.style.name if para.style else ''
        if 'Heading' in style_name or 'heading' in style_name:
            level = 0
            try:
                level = int(re.search(r'\d+', style_name).group())
            except:
                level = 0
            headings.append({
                'level': level,
                'style': style_name,
                'text': para.text.strip()[:120]
            })

    print(f"\n  📝 Tổng quan:")
    print(f"     Tổng paragraphs: {para_count}")
    print(f"     Tổng ký tự: {total_text_len:,}")
    print(f"     Ước tính trang: ~{total_text_len // 2000} trang")
    print(f"     Tổng headings: {len(headings)}")

    # 2. Phân tích heading levels
    level_counts = {}
    for h in headings:
        lv = h['level']
        level_counts[lv] = level_counts.get(lv, 0) + 1

    print(f"\n  📊 Phân bố Heading levels:")
    for lv in sorted(level_counts.keys()):
        print(f"     Heading {lv}: {level_counts[lv]} mục")

    # 3. Tables
    table_count = len(doc.tables)
    print(f"\n  📋 Bảng (Tables): {table_count}")
    for i, table in enumerate(doc.tables):
        rows = len(table.rows)
        cols = len(table.columns)
        # Lấy header row
        header = ''
        if rows > 0:
            cells = [cell.text.strip()[:30] for cell in table.rows[0].cells]
            header = ' | '.join(cells)
        print(f"     Bảng {i+1}: {rows}×{cols} — {header[:80]}")

    # 4. Images
    image_count = 0
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            image_count += 1
    print(f"\n  🖼️  Hình ảnh: {image_count}")

    # 5. Liệt kê headings chi tiết
    print(f"\n  📑 Danh mục Heading đầy đủ:")
    print(f"  {'─'*75}")
    for h in headings:
        indent = '    ' * h['level']
        text = h['text'] if h['text'] else '(trống)'
        print(f"     {indent}[H{h['level']}] {text}")

    return {
        'label': label,
        'para_count': para_count,
        'total_chars': total_text_len,
        'est_pages': total_text_len // 2000,
        'headings': headings,
        'heading_count': len(headings),
        'level_counts': level_counts,
        'table_count': table_count,
        'image_count': image_count,
    }


def compare_reports(data_tuan, data_minh):
    """So sánh 2 báo cáo."""
    print(f"\n\n{'='*80}")
    print(f"  SO SÁNH 2 BÁO CÁO")
    print(f"{'='*80}")

    metrics = [
        ('Tổng paragraphs', 'para_count'),
        ('Tổng ký tự', 'total_chars'),
        ('Ước tính trang', 'est_pages'),
        ('Số headings', 'heading_count'),
        ('Số bảng', 'table_count'),
        ('Số hình ảnh', 'image_count'),
    ]

    print(f"\n  {'Chỉ số':<25} {'Tuân':>12} {'Minh':>12} {'Chênh lệch':>12}")
    print(f"  {'─'*65}")
    for label, key in metrics:
        v_t = data_tuan[key]
        v_m = data_minh[key]
        diff = v_m - v_t
        sign = '+' if diff > 0 else ''
        print(f"  {label:<25} {v_t:>12,} {v_m:>12,} {sign}{diff:>11,}")

    # So sánh chapter structure
    print(f"\n  📌 Nhận xét:")

    if data_minh['est_pages'] < data_tuan['est_pages'] * 0.5:
        print(f"     ⚠️  Báo cáo Minh ngắn hơn nhiều so với Tuân ({data_minh['est_pages']} vs {data_tuan['est_pages']} trang)")

    if data_minh['table_count'] < data_tuan['table_count']:
        print(f"     ⚠️  Minh thiếu bảng ({data_minh['table_count']} vs {data_tuan['table_count']})")

    if data_minh['image_count'] < data_tuan['image_count']:
        print(f"     ⚠️  Minh thiếu hình ảnh ({data_minh['image_count']} vs {data_tuan['image_count']})")

    # Heading keywords from Tuân that Minh might be missing
    tuan_headings_text = set(h['text'].lower().strip() for h in data_tuan['headings'] if h['text'].strip())
    minh_headings_text = set(h['text'].lower().strip() for h in data_minh['headings'] if h['text'].strip())

    missing = tuan_headings_text - minh_headings_text
    if missing:
        print(f"\n  📝 Headings có ở Tuân nhưng Minh KHÔNG có (mẫu):")
        for m in sorted(list(missing))[:20]:
            print(f"     - {m[:80]}")


def main():
    results = {}
    for key, filepath in FILES.items():
        if filepath.exists():
            label = 'Nguyễn Thanh Tuân' if key == 'tuan' else 'Tăng Tuấn Minh'
            results[key] = analyze_docx(str(filepath), label)
        else:
            print(f"⚠️  File không tồn tại: {filepath}")

    if len(results) == 2:
        compare_reports(results['tuan'], results['minh'])

    # Save summary as JSON
    output_path = BASE_DIR / 'report_analysis.json'
    # Convert for JSON serialization
    json_data = {}
    for key, data in results.items():
        json_data[key] = {
            'label': data['label'],
            'para_count': data['para_count'],
            'total_chars': data['total_chars'],
            'est_pages': data['est_pages'],
            'heading_count': data['heading_count'],
            'table_count': data['table_count'],
            'image_count': data['image_count'],
            'level_counts': data['level_counts'],
            'headings': data['headings']
        }

    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(json_data, f, indent=2, ensure_ascii=False)
    print(f"\n  💾 Saved analysis to: {output_path}")


if __name__ == '__main__':
    main()
