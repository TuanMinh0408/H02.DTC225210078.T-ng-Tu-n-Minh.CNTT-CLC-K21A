from docx import Document

def main():
    doc = Document('DTC225210134_Nguyễn Thanh Tuân_CNTTK21CLC.docx')
    print("TOTAL TABLES IN ORIGINAL:", len(doc.tables))
    with open('tuan_tables.txt', 'w', encoding='utf-8') as f:
        for idx, table in enumerate(doc.tables):
            f.write(f"\n========================================\n")
            f.write(f"TABLE {idx+1} ({len(table.rows)}x{len(table.columns)}):\n")
            f.write(f"========================================\n")
            for r_idx, row in enumerate(table.rows):
                row_text = []
                for c_idx, cell in enumerate(row.cells):
                    row_text.append(f"[C{c_idx}]: {cell.text.strip().replace(chr(10), ' ')}")
                f.write(f"Row {r_idx}: {' | '.join(row_text)}\n")

if __name__ == '__main__':
    main()
