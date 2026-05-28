from docx import Document
from docx.oxml.ns import qn

def inspect_sections():
    doc = Document("DTC225210134_Nguyễn Thanh Tuân_CNTTK21CLC.docx")
    print(f"Total paragraphs: {len(doc.paragraphs)}")
    print(f"Total sections: {len(doc.sections)}")
    
    # Let's find which paragraphs contain section breaks
    # Section breaks are stored in the paragraph's xml: <w:pPr><w:sectPr>...
    p_breaks = []
    for i, p in enumerate(doc.paragraphs):
        pPr = p._p.get_or_add_pPr()
        sectPr = pPr.find(qn('w:sectPr'))
        if sectPr is not None:
            p_breaks.append((i, p.text.strip()))
            
    print("\nParagraphs with section breaks (at end of section):")
    for idx, text in p_breaks:
        print(f"Paragraph [{idx}]: {text}")
        
    print("\nSection details:")
    for idx, section in enumerate(doc.sections):
        print(f"\nSection {idx}:")
        print(f"  Start type: {section.start_type}")
        print(f"  Page Width: {section.page_width.cm:.2f} cm")
        print(f"  Page Height: {section.page_height.cm:.2f} cm")
        print(f"  Top Margin: {section.top_margin.cm:.2f} cm")
        print(f"  Bottom Margin: {section.bottom_margin.cm:.2f} cm")
        print(f"  Left Margin: {section.left_margin.cm:.2f} cm")
        print(f"  Right Margin: {section.right_margin.cm:.2f} cm")
        
        # Headers/Footers
        print(f"  Has Header: {section.header is not None and len(section.header.paragraphs[0].text.strip()) > 0}")
        if section.header:
            print(f"    Header Text: '{section.header.paragraphs[0].text.strip()}'")
        print(f"  Has Footer: {section.footer is not None and len(section.footer.paragraphs[0].text.strip()) > 0}")
        if section.footer:
            print(f"    Footer Text: '{section.footer.paragraphs[0].text.strip()}'")
            
if __name__ == "__main__":
    inspect_sections()
