from docx import Document

def main():
    doc = Document("DTC225210078_Tăng Tuấn Minh_CNTTK21CLC.docx")
    keywords = ['đề cương', 'clo', 'nhiệm vụ', 'tiến độ', 'tuần', 'lãnh đạo bộ môn']
    
    print("Searching in paragraphs:")
    for idx, p in enumerate(doc.paragraphs):
        text = p.text.lower()
        for kw in keywords:
            if kw in text:
                print(f"P [{idx}]: {repr(p.text[:100])} (Matched: {kw})")
                
    print("\nSearching in tables:")
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                text = cell.text.lower()
                for kw in keywords:
                    if kw in text:
                        print(f"Table {t_idx+1}, Row {r_idx}, Col {c_idx}: {repr(cell.text[:100])} (Matched: {kw})")

if __name__ == '__main__':
    main()
