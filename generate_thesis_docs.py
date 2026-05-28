import urllib.request
import urllib.parse
from docx import Document
from docx.shared import Inches
import os

def download_image(dot_string, filename):
    encoded = urllib.parse.quote(dot_string)
    url = f"https://quickchart.io/graphviz?graph={encoded}&format=png"
    try:
        urllib.request.urlretrieve(url, filename)
        return True
    except Exception as e:
        print(f"Error downloading {filename}: {e}")
        return False

# Hinh 2.1: Flow (Horizontal)
dot1 = """digraph G {
    rankdir=LR;
    node [shape=box, style="rounded,filled", fillcolor=white, color=black, fontname="Arial", fontsize=13, penwidth=1.5, margin=0.2];
    edge [color=black, fontname="Arial", fontsize=11, penwidth=1.5];
    
    A [label="Video / Camera"];
    B [label="YOLOv11-Pose\\n(Cắt vùng người)"];
    C [label="MediaPipe BlazePose\\n(Trích xuất 33 khớp)"];
    D [label="Bi-GRU Model\\n(Phân tích chuỗi)"];
    E [label="Kết quả\\n(Té ngã / Bình thường)", shape=ellipse];
    
    A -> B;
    B -> C;
    C -> D;
    D -> E;
}"""

# Hinh 2.2: Layered Architecture using HTML table
dot2 = """digraph G {
    node [shape=plaintext, fontname="Arial"];
    
    Architecture [label=<
    <TABLE BORDER="0" CELLBORDER="1" CELLSPACING="10" CELLPADDING="15">
      <TR><TD BORDER="0"><FONT POINT-SIZE="16"><B>Lớp Giao diện (Presentation Layer)</B></FONT></TD></TR>
      <TR><TD BORDER="0"><TABLE BORDER="0" CELLSPACING="15"><TR>
        <TD BORDER="2" STYLE="ROUNDED" WIDTH="180"><FONT POINT-SIZE="14">Web Dashboard</FONT></TD>
        <TD BORDER="2" STYLE="ROUNDED" WIDTH="180"><FONT POINT-SIZE="14">Hệ thống Cảnh báo</FONT></TD>
      </TR></TABLE></TD></TR>
      
      <TR><TD BORDER="0"><FONT POINT-SIZE="16"><B>Lớp Dịch vụ (Service Layer)</B></FONT></TD></TR>
      <TR><TD BORDER="0"><TABLE BORDER="0" CELLSPACING="15"><TR>
        <TD BORDER="2" STYLE="ROUNDED" WIDTH="180"><FONT POINT-SIZE="14">FastAPI Backend</FONT></TD>
        <TD BORDER="2" STYLE="ROUNDED" WIDTH="180"><FONT POINT-SIZE="14">WebSocket Server</FONT></TD>
      </TR></TABLE></TD></TR>

      <TR><TD BORDER="0"><FONT POINT-SIZE="16"><B>Lớp AI &amp; Xử lý (AI Processing Layer)</B></FONT></TD></TR>
      <TR><TD BORDER="0"><TABLE BORDER="0" CELLSPACING="15"><TR>
        <TD BORDER="2" STYLE="ROUNDED" WIDTH="150"><FONT POINT-SIZE="14">YOLOv11-Pose</FONT></TD>
        <TD BORDER="2" STYLE="ROUNDED" WIDTH="150"><FONT POINT-SIZE="14">MediaPipe</FONT></TD>
        <TD BORDER="2" STYLE="ROUNDED" WIDTH="150"><FONT POINT-SIZE="14">Bi-GRU Model</FONT></TD>
      </TR></TABLE></TD></TR>

      <TR><TD BORDER="0"><FONT POINT-SIZE="16"><B>Lớp Dữ liệu (Data Layer)</B></FONT></TD></TR>
      <TR><TD BORDER="0"><TABLE BORDER="0" CELLSPACING="15"><TR>
        <TD BORDER="2" STYLE="ROUNDED" WIDTH="180"><FONT POINT-SIZE="14">Camera / Webcam</FONT></TD>
        <TD BORDER="2" STYLE="ROUNDED" WIDTH="180"><FONT POINT-SIZE="14">SisFall Dataset</FONT></TD>
      </TR></TABLE></TD></TR>
    </TABLE>
    >];
}"""

# Hinh 2.3: Flowchart (Vertical, Orthogonal)
dot3 = """digraph G {
    rankdir=TB;
    splines=ortho;
    nodesep=0.6;
    ranksep=0.5;
    node [style="filled", fillcolor=white, color=black, fontname="Arial", fontsize=12, penwidth=1.5];
    edge [color=black, fontname="Arial", fontsize=11, penwidth=1.5];
    
    Start [shape=oval, label="BẮT ĐẦU", penwidth=2, margin=0.1];
    Read [shape=box, label="Đọc khung hình\\ntừ Camera", margin=0.2];
    YOLO [shape=box, label="YOLOv11-Pose\\nPhát hiện người", margin=0.2];
    D1 [shape=diamond, label="Có phát hiện\\nngười?", margin=0.1];
    MP [shape=box, label="MediaPipe\\nTrích 33 khớp", margin=0.2];
    Buffer [shape=box, label="Cập nhật bộ đệm\\nchuỗi thời gian", margin=0.2];
    D2 [shape=diamond, label="Đủ chuỗi\\n30 frame?", margin=0.1];
    BiGRU [shape=box, label="Mô hình Bi-GRU\\nDự đoán hành vi", margin=0.2];
    D3 [shape=diamond, label="Là hành vi\\nTé ngã?", margin=0.1];
    Alert [shape=box, label="Gửi cảnh báo\\nlên Dashboard", margin=0.2];
    
    Start -> Read;
    Read -> YOLO;
    YOLO -> D1;
    D1 -> MP [label="Có"];
    
    MP -> Buffer;
    Buffer -> D2;
    D2 -> BiGRU [label="Có"];
    
    BiGRU -> D3;
    D3 -> Alert [label="Có"];
    
    // Loops routed to the left
    D1 -> Read [label="Không", tailport=w, headport=w];
    D2 -> Read [label="Không", tailport=w, headport=w];
    D3 -> Read [label="Không", tailport=w, headport=w];
    
    // Loop back from alert routed to the right
    Alert -> Read [tailport=e, headport=e];
}"""

print("Downloading images via API...")
s1 = download_image(dot1, "img_2_1_pro.png")
s2 = download_image(dot2, "img_2_2_pro.png")
s3 = download_image(dot3, "img_2_3_pro.png")

if s1 and s2 and s3:
    print("Creating docx...")
    doc = Document()
    doc.add_heading('Sơ đồ Chương 2 (Chuẩn Báo cáo Đồ án)', 0)

    p1 = doc.add_paragraph()
    r1 = p1.add_run('Hình 2.1: Sơ đồ luồng kết hợp YOLO và trích xuất khung xương')
    r1.bold = True
    doc.add_picture("img_2_1_pro.png", width=Inches(6.0))

    p2 = doc.add_paragraph()
    r2 = p2.add_run('\nHình 2.2: Sơ đồ kiến trúc tổng thể hệ thống')
    r2.bold = True
    doc.add_picture("img_2_2_pro.png", width=Inches(6.0))

    p3 = doc.add_paragraph()
    r3 = p3.add_run('\nHình 2.3: Lưu đồ thuật toán quy trình phát hiện té ngã')
    r3.bold = True
    doc.add_picture("img_2_3_pro.png", width=Inches(5.0))

    save_path = r'C:\Users\minhc\.gemini\antigravity\scratch\graduation_project\Sodo_Chuong2_Chuan.docx'
    doc.save(save_path)
    print(f"Done. Saved to {save_path}")
else:
    print("Failed to download images.")
