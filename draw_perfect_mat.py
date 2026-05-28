import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Ellipse, Polygon
import matplotlib.path as mpath
import matplotlib.patches as mpatches
import os

artifact_dir = r"C:\Users\minhc\.gemini\antigravity\scratch\graduation_project"
os.chdir(artifact_dir)

def draw_2_2():
    fig, ax = plt.subplots(figsize=(15, 6))
    ax.set_xlim(0, 150)
    ax.set_ylim(0, 60)
    ax.axis('off')

    layers = [
        {"x": 5, "w": 30, "title": "Lớp Dữ liệu", "nodes": ["Camera / Webcam", "SisFall Dataset"]},
        {"x": 40, "w": 35, "title": "Lớp AI & Xử lý", "nodes": ["YOLOv11-Pose", "MediaPipe BlazePose", "Bi-GRU Model"]},
        {"x": 80, "w": 30, "title": "Lớp Dịch vụ", "nodes": ["FastAPI Backend", "WebSocket Server"]},
        {"x": 115, "w": 30, "title": "Lớp Giao diện", "nodes": ["Web Dashboard", "Hệ thống Cảnh báo"]}
    ]

    def add_ellipse(x, y, w, h, text):
        ell = Ellipse((x, y), w, h, ec="black", fc="white", lw=2)
        ax.add_patch(ell)
        lines = text.split('\n')
        if len(lines) == 1:
            ax.text(x, y, text, ha='center', va='center', fontsize=11, fontweight='bold', color='black')
        else:
            ax.text(x, y+1.5, lines[0], ha='center', va='center', fontsize=11, fontweight='bold', color='black')
            ax.text(x, y-2, lines[1], ha='center', va='center', fontsize=10, color='black')

    def add_arrow(x1, y1, x2, y2):
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                    arrowprops=dict(arrowstyle="->", lw=2, color="black"))

    for L in layers:
        x = L["x"]
        w = L["w"]
        box = FancyBboxPatch((x, 5), w, 45, boxstyle="round,pad=1", ec="black", fc="#f8f9fa", lw=1.5, ls="--")
        ax.add_patch(box)
        ax.text(x + w/2, 53, L["title"], ha='center', va='center', fontsize=13, fontweight='bold')
        
        n = len(L["nodes"])
        spacing = 40 / n
        for i, node in enumerate(reversed(L["nodes"])):
            y = 5 + i * spacing + spacing/2
            if " " in node and len(node) > 15 and "\n" not in node:
                words = node.split(" ")
                mid = len(words)//2
                node = " ".join(words[:mid]) + "\n" + " ".join(words[mid:])
            add_ellipse(x + w/2, y, w-4, 9, node)

    # Arrows between layers
    add_arrow(35, 27.5, 40, 27.5)
    add_arrow(75, 27.5, 80, 27.5)
    add_arrow(110, 27.5, 115, 27.5)

    plt.savefig("img_2_2_mat.png", dpi=300, bbox_inches='tight', facecolor='white')
    plt.close()

def draw_2_3():
    fig, ax = plt.subplots(figsize=(20, 7))
    ax.set_xlim(0, 240)
    ax.set_ylim(0, 70)
    ax.axis('off')

    def add_ellipse(x, y, w, h, text):
        ell = Ellipse((x, y), w, h, ec="black", fc="white", lw=2)
        ax.add_patch(ell)
        lines = text.split('\n')
        if len(lines) == 1:
            ax.text(x, y, text, ha='center', va='center', fontsize=11, fontweight='bold', color='black')
        else:
            ax.text(x, y+1.5, lines[0], ha='center', va='center', fontsize=11, fontweight='bold', color='black')
            ax.text(x, y-2, lines[1], ha='center', va='center', fontsize=10, color='black')

    def add_diamond(x, y, w, h, text):
        pts = [[x, y+h/2], [x+w/2, y+h], [x+w, y+h/2], [x+w/2, y]]
        poly = Polygon(pts, closed=True, ec="black", fc="white", lw=2)
        ax.add_patch(poly)
        lines = text.split('\n')
        if len(lines) == 1:
            ax.text(x+w/2, y+h/2, text, ha='center', va='center', fontsize=10, fontweight='bold', color='black')
        else:
            ax.text(x+w/2, y+h/2+1.5, lines[0], ha='center', va='center', fontsize=10, fontweight='bold', color='black')
            ax.text(x+w/2, y+h/2-1.5, lines[1], ha='center', va='center', fontsize=10, color='black')

    def add_arrow(x1, y1, x2, y2, label=""):
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                    arrowprops=dict(arrowstyle="->", lw=2, color="black"))
        if label:
            ax.text((x1+x2)/2, (y1+y2)/2 + 2.5, label, ha='center', va='center', fontsize=11, fontweight='bold')

    def add_path(pts, label=""):
        path_data = [mpath.Path.MOVETO] + [mpath.Path.LINETO]*(len(pts)-1)
        path = mpath.Path(pts, path_data)
        patch = mpatches.PathPatch(path, fill=False, lw=2, color="black")
        ax.add_patch(patch)
        ax.annotate("", xy=pts[-1], xytext=pts[-2], arrowprops=dict(arrowstyle="->", lw=2, color="black"))
        if label:
            ax.text(pts[1][0]+2, pts[1][1]-2, label, ha='left', va='center', fontsize=11, fontweight='bold')

    cy = 35
    w_e = 16
    h_e = 10
    
    # Coordinates
    x_start = 10
    x_read = 34
    x_yolo = 58
    x_d1 = 78
    x_mp = 104
    x_buf = 128
    x_d2 = 148
    x_gru = 174
    x_d3 = 194
    x_alert = 220

    # Start
    add_ellipse(x_start, cy, w_e, h_e, "BẮT ĐẦU")
    add_arrow(x_start+w_e/2, cy, x_read-w_e/2, cy)

    add_ellipse(x_read, cy, w_e, h_e, "Đọc frame\ntừ Camera")
    add_arrow(x_read+w_e/2, cy, x_yolo-w_e/2, cy)

    add_ellipse(x_yolo, cy, w_e, h_e, "YOLOv11\nPhát hiện")
    add_arrow(x_yolo+w_e/2, cy, x_d1, cy)

    # D1
    add_diamond(x_d1, cy-7, 14, 14, "Có người?")
    add_arrow(x_d1+14, cy, x_mp-w_e/2, cy, "Có")
    # Loop D1 UP
    add_path([(x_d1+7, cy+7), (x_d1+7, 50), (x_read, 50), (x_read, cy+h_e/2)], "Không")

    # MP
    add_ellipse(x_mp, cy, w_e, h_e, "MediaPipe\nTrích khớp")
    add_arrow(x_mp+w_e/2, cy, x_buf-w_e/2, cy)

    # Buffer
    add_ellipse(x_buf, cy, w_e, h_e, "Cập nhật\nbộ đệm")
    add_arrow(x_buf+w_e/2, cy, x_d2, cy)

    # D2
    add_diamond(x_d2, cy-7, 14, 14, "Đủ 30\nframe?")
    add_arrow(x_d2+14, cy, x_gru-w_e/2, cy, "Có")
    # Loop D2 UP higher
    add_path([(x_d2+7, cy+7), (x_d2+7, 56), (x_read-3, 56), (x_read-3, cy+h_e/2)], "Không")

    # Bi-GRU
    add_ellipse(x_gru, cy, w_e, h_e, "Bi-GRU\nDự đoán")
    add_arrow(x_gru+w_e/2, cy, x_d3, cy)

    # D3
    add_diamond(x_d3, cy-7, 14, 14, "Té ngã?")
    add_arrow(x_d3+14, cy, x_alert-w_e/2, cy, "Có")
    # Loop D3 UP highest
    add_path([(x_d3+7, cy+7), (x_d3+7, 62), (x_read-6, 62), (x_read-6, cy+h_e/2)], "Không")

    # Alert
    add_ellipse(x_alert, cy, w_e, h_e, "Cảnh báo\nDashboard")
    # Loop Alert DOWN
    add_path([(x_alert, cy-h_e/2), (x_alert, 15), (x_read+3, 15), (x_read+3, cy-h_e/2)], "")

    plt.savefig("img_2_3_mat.png", dpi=300, bbox_inches='tight', facecolor='white')
    plt.close()

print("Drawing 2.2 and 2.3 with Matplotlib...")
draw_2_2()
draw_2_3()

from docx import Document
from docx.shared import Inches
doc = Document()
doc.add_heading('Sơ đồ Chương 2 (Bản Tối Ưu Tùy Chỉnh Cực Nét)', 0)

doc.add_paragraph('Hình 2.2: Sơ đồ kiến trúc tổng thể hệ thống (Xếp ngang, Khối Use Case)')
doc.add_picture("img_2_2_mat.png", width=Inches(6.0))

doc.add_paragraph('Hình 2.3: Lưu đồ thuật toán quy trình phát hiện té ngã (Xếp ngang, Căn lề chính xác 100%)')
doc.add_picture("img_2_3_mat.png", width=Inches(6.5))

save_path = 'Sodo_Chuong2_Final.docx'
doc.save(save_path)
print(f"Done. Saved to {save_path}")
