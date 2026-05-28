from docx import Document

def main():
    doc = Document('DTC225210078_Tăng Tuấn Minh_CNTTK21CLC.docx')
    print("Searching in Minh's generated doc...")
    keywords = ['đề cương', 'clo', 'nhiệm vụ', 'tiến độ', 'lãnh đạo bộ môn']
    for i, p in enumerate(doc.paragraphs):
        text_lower = p.text.lower()
        for kw in keywords:
            if kw in text_lower:
                print(f"Paragraph [{i:04d}]: '{p.text}'")
                break
                
    for i, t in enumerate(doc.tables):
        for r_idx, row in enumerate(t.rows):
            for c_idx, cell in enumerate(row.cells):
                cell_lower = cell.text.lower()
                for kw in keywords:
                    if kw in cell_lower:
                        print(f"Table [{i}][Row {r_idx}][Col {c_idx}]: {repr(cell.text[:100])}")
                        break

if __name__ == '__main__':
    main()
