from docx import Document
import sys

def main():
    original_path = r"C:\Users\minhc\Downloads\Đồ án tốt nghiệp\H02.DTC225210078.Tăng Tuấn Minh.CNTTK21CLC.docx"
    try:
        doc = Document(original_path)
        print("="*60)
        print(f"Original Doc: {original_path}")
        print(f"Total Paragraphs: {len(doc.paragraphs)}")
        print(f"Total Tables: {len(doc.tables)}")
        print(f"Total Sections: {len(doc.sections)}")
        print("="*60)
        
        print("\nFIRST 40 PARAGRAPHS:")
        for i, p in enumerate(doc.paragraphs[:40]):
            style_info = p.style.name if p.style else "NoStyle"
            text_preview = p.text.strip()[:80]
            print(f"[{i:02d}] Style: {style_info:<20} | Text: {text_preview}")
            
        print("\nTABLE DETAILS:")
        for i, t in enumerate(doc.tables):
            rows = len(t.rows)
            cols = len(t.columns)
            print(f"Table {i+1}: {rows}x{cols}")
            if rows > 0:
                cells = [c.text.strip()[:20] for c in t.rows[0].cells]
                print(f"   Header: {' | '.join(cells)}")
                
    except Exception as e:
        print(f"Error inspecting file: {e}")

if __name__ == '__main__':
    main()
