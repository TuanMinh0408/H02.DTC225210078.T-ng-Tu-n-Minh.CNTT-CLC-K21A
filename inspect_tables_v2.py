from docx import Document

def main():
    doc = Document('DTC225210134_Nguyễn Thanh Tuân_CNTTK21CLC.docx')
    body = doc.element.body
    
    with open('tuan_tables_info.txt', 'w', encoding='utf-8') as f:
        table_idx = 0
        p_idx = 0
        for child in body:
            if child.tag.endswith('p'):
                p_obj = [p for p in doc.paragraphs if p._element == child]
                text = p_obj[0].text if p_obj else ""
                if text.strip():
                    f.write(f"P {p_idx}: {text[:100]}\n")
                p_idx += 1
            elif child.tag.endswith('tbl'):
                t_obj = [t for t in doc.tables if t._element == child]
                if t_obj:
                    t = t_obj[0]
                    first_cell = t.rows[0].cells[0].text.strip() if t.rows and t.rows[0].cells else ""
                    f.write(f"--- TABLE {table_idx} at body index {body.index(child)}, size: {len(t.rows)}x{len(t.columns)}, first cell: {first_cell[:50]} ---\n")
                table_idx += 1

if __name__ == '__main__':
    main()
