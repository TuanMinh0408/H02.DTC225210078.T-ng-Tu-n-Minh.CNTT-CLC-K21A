from docx import Document
import io
import os

docx_path = r'C:\Users\minhc\Downloads\Đồ án tốt nghiệp\H02.DTC225210078.Tăng Tuấn Minh.CNTTK21CLC.docx'
output_path = 'outline_extracted_v2.txt'

try:
    doc = Document(docx_path)
    full_content = []
    
    for element in doc.element.body.xpath('.//w:p | .//w:tbl'):
        if element.tag.endswith('p'):
            para = [p for p in doc.paragraphs if p._element == element]
            if para:
                full_content.append(para[0].text)
        elif element.tag.endswith('tbl'):
            table = [t for t in doc.tables if t._element == element][0]
            for row in table.rows:
                row_text = ' | '.join([cell.text.strip().replace('\n', ' ') for cell in row.cells])
                full_content.append(row_text)
                
    text = '\n'.join(full_content)
    with io.open(output_path, 'w', encoding='utf-8') as f:
        f.write(text)
    print(f"Successfully extracted text and tables to {output_path}")
except Exception as e:
    print(f"Error: {e}")
