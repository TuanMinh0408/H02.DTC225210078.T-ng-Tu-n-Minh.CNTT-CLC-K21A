from docx import Document

def inspect_doc(filepath, label):
    print(f"\n==================== {label} ====================")
    try:
        doc = Document(filepath)
        print("Total Paragraphs:", len(doc.paragraphs))
        print("Total Sections:", len(doc.sections))
        print("Total Tables:", len(doc.tables))
        
        print("\nHeadings:")
        h_count = 0
        for i, p in enumerate(doc.paragraphs):
            style_name = p.style.name if p.style else ""
            text = p.text.strip()
            # If paragraph has any heading style or is all uppercase and short, print it
            if 'Heading' in style_name or (text.isupper() and len(text) < 100 and len(text) > 3):
                h_count += 1
                print(f"[{i:04d}] Style: {style_name:<15} | Text: {repr(text)}")
        print("Total Headings printed:", h_count)
    except Exception as e:
        print("Error:", e)

def main():
    inspect_doc("DTC225210134_Nguyễn Thanh Tuân_CNTTK21CLC.docx", "Nguyễn Thanh Tuân")
    inspect_doc("DTC225210078_Tăng Tuấn Minh_CNTTK21CLC.docx", "Tăng Tuấn Minh (Generated)")

if __name__ == '__main__':
    main()
