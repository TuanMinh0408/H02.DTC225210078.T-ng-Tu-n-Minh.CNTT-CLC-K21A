from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_border(cell, **kwargs):
    """
    Hàm hỗ trợ kẻ bảng chuẩn (Grid)
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for edge in ('top', 'start', 'bottom', 'end'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = tcPr.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcPr.append(element)
            for key, value in edge_data.items():
                element.set(qn('w:{}'.format(key)), str(value))

def create_scientific_report():
    doc = Document()
    
    # Font settings
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(13)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('BÁO CÁO TIẾN ĐỘ THỰC HIỆN ĐỀ TÀI TỐT NGHIỆP')
    run.bold = True
    run.font.size = Pt(16)
    run.font.name = 'Times New Roman'

    # Subject Info
    p = doc.add_paragraph()
    p.add_run('Tên đề tài: ').bold = True
    p.add_run('Ứng dụng học sâu trong nhận dạng hành vi té ngã của người cao tuổi')
    
    info_lines = [
        ('Sinh viên thực hiện', 'Tăng Tuấn Minh'),
        ('Mã số sinh viên', 'DTC225210078'),
        ('Lớp / Khóa', 'CNTT K21CLC'),
        ('Giáo viên hướng dẫn', 'TS. Nguyễn Tuấn Anh'),
        ('Thời điểm báo cáo', '16/04/2026 (Tuần thứ 12 - Giai đoạn 1)'),
    ]
    for label, val in info_lines:
        p = doc.add_paragraph()
        p.add_run(f'{label}: ').bold = True
        p.add_run(val)

    # I. THÔNG TIN CHUNG
    doc.add_heading('I. THÔNG TIN CHUNG VỀ ĐỀ TÀI', level=1)
    doc.add_paragraph('1.1. Mục tiêu công nghệ', style='List Bullet')
    doc.add_paragraph('Nghiên cứu và triển khai hệ thống AI có khả năng nhận diện hành vi té ngã của người cao tuổi với độ chính xác cao (>95%) thông qua việc tổ hợp đặc trưng (Feature Fusion) từ hai nguồn: (1) Chuỗi tọa độ khung xương 17 điểm Pose trích xuất từ Video và (2) Tín hiệu gia tốc tuyến tính từ cảm biến wearable.')
    
    # II. TIẾN ĐỘ CHI TIẾT
    doc.add_heading('II. TIẾN ĐỘ THỰC HIỆN ĐẾN THỜI ĐIỂM HIỆN TẠI', level=1)
    doc.add_paragraph('Tính đến ngày báo cáo, dự án đã hoàn thành toàn bộ khối lượng công việc của 12 tuần triển khai theo đúng đề cương dự án.')
    
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    cols = ['Hạng mục', 'Kế hoạch đề cương', 'Thực tế triển khai', 'Đánh giá']
    for i, name in enumerate(cols):
        hdr[i].text = name
        hdr[i].paragraphs[0].runs[0].bold = True

    data = [
        ['Khảo sát & SRS', 'Tuần 1 - 3', 'Hoàn thành báo cáo khảo sát và kiến trúc hệ thống tổng quát.', 'Đúng tiến độ'],
        ['Dashboard & Pipe', 'Tuần 4 - 6', 'Xây dựng Dashboard giám sát real-time (FPS > 25), tích hợp WebSocket.', 'Vượt tiến độ'],
        ['Vision Engine', 'Tuần 7 - 9', 'Tối ưu hóa YOLOv11-Pose để trích xuất tọa độ 17 điểm khung xương thời gian thực.', 'Đúng tiến độ'],
        ['Data Pipeline', 'Tuần 10 - 12', 'Xây dựng Data Loader đa phương thức, đồng bộ hóa CSV cảm biến và khung hình.', 'Vượt tiến độ'],
        ['LSTM Core', 'Tuần 12', 'Hoàn thiện kiến trúc mạng Fusion LSTM nhận đầu vào 40 tham số.', 'Hoàn thành sớm']
    ]
    for row_data in data:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = val

    # III. THỐNG KÊ CHI TIẾT
    doc.add_heading('III. THỐNG KÊ SẢN PHẨM DỰ ÁN', level=1)
    stats_table = doc.add_table(rows=1, cols=2)
    stats_table.style = 'Table Grid'
    stats_data = [
        ['Tổng số mã nguồn (Python)', '~2.500 dòng code'],
        ['Số lượng Models tích hợp', '02 Models (YOLOv11-Pose, FusionLSTM)'],
        ['Dữ liệu xử lý', '02 Video thực nghiệm + 195 samples cảm biến UP-Fall'],
        ['Cơ sở dữ liệu', 'SQLite 3 (Lưu trữ lịch sử sự cố & Metadata)'],
        ['Giao diện', '01 Dashboard nền tảng Web (HTML5/Vanilla CSS/JS)']
    ]
    for label, val in stats_data:
        row = stats_table.add_row().cells
        row[0].text = label
        row[1].text = val

    # IV. KHÓ KHĂN VÀ GIẢI PHÁP KỸ THUẬT
    doc.add_heading('IV. KHÓ KHĂN VÀ CÁC GIẢI PHÁP ĐÃ ÁP DỤNG', level=1)
    diffs = [
        ('Độ lệch tần số lấy mẫu (Sampling Rate Mismatch)', 'Tần số camera (30Hz) không khớp với cảm biến (đến 100Hz). Giải pháp: Xây dựng module Sync dựa trên kỹ thuật nội suy (Interpolation) để đồng bộ hóa vector đặc trưng.'),
        ('Hiện tượng che khuất (Occlusion)', 'Camera thỉnh thoảng mất dấu keypoints do góc quay. Giải pháp: Sử dụng dữ liệu cảm biến (Inertial data) làm nguồn bù đắp tin cậy trong các khung hình lỗi.'),
        ('Độ trễ truyền tải (Transmission Latency)', 'Truyền Base64 qua WebSocket gây lag. Giải pháp: Tối ưu hóa kích thước frame hình và sử dụng cơ chế xử lý bất đồng bộ (Asynchronous processing).')
    ]
    for title, sol in diffs:
        p = doc.add_paragraph()
        p.add_run(f'+ {title}: ').bold = True
        p.add_run(sol)

    # V. KẾ HOẠCH GIAI ĐOẠN TIẾP THEO
    doc.add_heading('V. KẾ HOẠCH GIAI ĐOẠN TIẾP THEO (TUẦN 13 - 18)', level=1)
    next_tasks = [
        'Tuần 13-14: Huấn luyện mô hình Fusion LSTM với đầy đủ tập dữ liệu UP-Fall (17 subjects).',
        'Tuần 15: Tinh chỉnh Hyperparameters và đánh giá độ chính xác (Precision/Recall) so với Vision-only.',
        'Tuần 16: Tối ưu hóa mô hình sang định dạng TensorRT hoặc CoreML để triển khai thiết bị đầu cuối.',
        'Tuần 17-18: Hoàn thiện báo cáo bản thảo và chuẩn bị video demo hệ thống cuối cùng.'
    ]
    for t in next_tasks:
        doc.add_paragraph(t, style='List Number')

    # VI. KẾ LUẬN
    doc.add_heading('VI. KẾ LUẬN', level=1)
    doc.add_paragraph('Sau Giai đoạn 1, đề tài đã xây dựng thành công nền tảng kỹ thuật cốt lõi. Việc tích hợp dữ liệu đa phương thức đã được chứng minh là khả thi qua các thử nghiệm ban đầu. Hệ thống hiện đang đi đúng lộ trình và dự kiến sẽ hoàn thành đúng hạn với chất lượng khoa học cao.')

    # Appendix
    doc.add_page_break()
    doc.add_heading('PHỤ LỤC A - DANH SÁCH API ENDPOINTS', level=1)
    api_table = doc.add_table(rows=1, cols=3)
    api_table.style = 'Table Grid'
    hdr = api_table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = 'Công cụ', 'Endpoint/Protocol', 'Chức năng'
    apis = [
        ['FASTAPI', 'GET /', 'Phục vụ Dashboard chính'],
        ['WEBSOCKET', 'WS /ws', 'Stream video & dự đoán AI'],
        ['API REST', 'GET /history', 'Truy vấn lịch sử sự cố'],
        ['STATIC', 'GET /captures/', 'Truy xuất file ảnh bằng chứng']
    ]
    for item in apis:
        row = api_table.add_row().cells
        for i, val in enumerate(item): row[i].text = val

    # Save
    save_path = 'BAO_CAO_TIEN_DO_TANG_TUAN_MINH.docx'
    doc.save(save_path)
    print(f"Final report generated: {save_path}")

if __name__ == "__main__":
    create_scientific_report()
